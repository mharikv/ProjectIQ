"""Generate the AI PM Agent functional report as a Word document."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUTPUT = r"C:\curser\aiAgent\AI_Project_Management_Agent_Report.docx"


def set_heading_style(doc):
    for i in range(1, 4):
        style = doc.styles[f"Heading {i}"]
        style.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
        style.font.bold = True


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = str(val)
    doc.add_paragraph()


def build():
    doc = Document()
    set_heading_style(doc)

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("AI Project Management Agent")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Comprehensive Functional Report")
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("\nApplication: ProjectIQ\n")
    meta.add_run("Prepared For: Project Management Professionals\n")
    meta.add_run("Document Type: Functional Overview & Business Value Analysis\n")
    meta.add_run("Date: August 2026\n")

    doc.add_page_break()

    # 1 Executive Summary
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "The AI Project Management Agent is a web-based application built specifically for manufacturing "
        "and engineering project managers. It brings together traditional project management disciplines—planning, "
        "scheduling, budgeting, resourcing, risk, quality, and maintenance—into a single workspace, enhanced by "
        "artificial intelligence."
    )
    doc.add_paragraph(
        "Rather than replacing the project manager, the application acts as an intelligent assistant: it helps draft "
        "plans faster, surfaces risks early, tracks performance in real time, and answers project questions on demand. "
        "Project managers can work through a structured dashboard, manage multiple projects at once, and move from "
        "high-level charter approval down to machine-level maintenance planning without switching tools."
    )
    doc.add_paragraph(
        "The application is especially suited to complex manufacturing programs—such as production line upgrades, "
        "new product development (e.g., electronics/hardware), and mass-production readiness projects—where hardware, "
        "software, supply chain, quality, and regulatory concerns must be managed together."
    )

    # 2 Purpose
    doc.add_heading("2. Purpose and Objectives", level=1)
    doc.add_heading("2.1 Purpose", level=2)
    doc.add_paragraph(
        "Manufacturing projects involve many interdependent workstreams: design, procurement, production, testing, "
        "certification, and go-live. Spreadsheets and disconnected documents often lead to outdated plans, missed risks, "
        "and poor visibility for stakeholders."
    )
    doc.add_paragraph("This application addresses that gap by providing:")
    for item in [
        "A central project hub for all planning and tracking data",
        "AI-assisted generation of charters, schedules, budgets, and risk registers",
        "Manual control so project managers can refine every detail",
        "A conversational assistant that answers questions and can update project data",
        "Exportable reports for reviews, approvals, and stakeholder communication",
    ]:
        add_bullet(doc, item)

    doc.add_heading("2.2 Objectives for Project Managers", level=2)
    add_table(doc, ["Objective", "How the Application Supports It"], [
        ["Faster project setup", "AI-generated charter, WBS, and schedule as starting points"],
        ["Better visibility", "Dashboard, KPIs, health score, and Gantt timeline"],
        ["Proactive risk management", "Risk register with scoring and mitigations"],
        ["Cost control", "Budget vs. actual tracking with variance and forecast"],
        ["Resource clarity", "Department-based resource allocation with utilization"],
        ["Operational readiness", "Inventory planning and predictive maintenance views"],
        ["Stakeholder reporting", "Downloadable PowerPoint project report"],
    ])

    # 3 Target Users
    doc.add_heading("3. Target Users", level=1)
    add_table(doc, ["User Role", "Primary Use"], [
        ["Project Manager", "Owns projects end-to-end; creates plans, tracks progress, reports to sponsors"],
        ["Administrator", "Oversees all projects; manages access and portfolio visibility"],
        ["Functional Leads", "Review and contribute to charter, WBS, risks, and resource plans"],
        ["Sponsors / Executives", "Review health, KPIs, and exported reports for go/no-go decisions"],
    ])

    # 4 Structure
    doc.add_heading("4. Application Structure at a Glance", level=1)
    doc.add_paragraph("The application is organized around three main areas:")
    for item in [
        "Login & Security — Secure access with role-based users",
        "Project Dashboard — Portfolio view of all projects",
        "Project Workspace — Detailed view of a single project with dedicated tabs for each management discipline",
    ]:
        add_bullet(doc, item)
    doc.add_paragraph(
        "Within each project, the project manager moves between Overview, Charter, WBS, Gantt, Resources, Budget, "
        "Inventory, Risks, Quality, Maintenance, KPIs, and AI Assistant—mirroring how manufacturing programs are "
        "actually run."
    )

    # 5 Detailed Functionalities
    doc.add_heading("5. Detailed Functionalities", level=1)

    sections = [
        ("5.1 User Authentication and Access Control",
         "Users sign in with a username and password. Different roles have appropriate access. Each project manager "
         "sees and manages their own projects; administrators can view the full portfolio.",
         ["Protects sensitive project and cost data", "Ensures accountability—each project has a clear owner",
          "Supports team-based deployment in a manufacturing organization"]),
        ("5.2 Project Dashboard and Portfolio Management",
         "The dashboard displays all projects showing name, status, priority, budget, health score, and target dates. "
         "Projects are organized into In Progress (planning, active, on hold) and Completed (finished or cancelled). "
         "Project managers can create new projects, open any project, and delete projects no longer needed.",
         ["Single place to see the entire portfolio", "Quick identification of troubled projects",
          "Clean separation between active and closed work", "Reduces time searching for project status"]),
        ("5.3 Project Overview and Core Details",
         "The Overview tab is the project control panel. Project managers can view and update name, description, status, "
         "priority, dates, budget, actual cost, budget utilization, and project health score.",
         ["Keeps fundamental metadata accurate", "Health score gives at-a-glance pulse check",
          "Status changes support portfolio reporting and phase-gate reviews"]),
        ("5.4 AI Project Charter Generator",
         "Generates a structured Project Charter including scope, objectives, deliverables, stakeholders, success criteria, "
         "and constraints/assumptions. Project managers can generate with AI, regenerate, or create and edit manually.",
         ["Accelerates project initiation", "Ensures consistency across projects",
          "Clear reference for scope disputes and change control", "Stakeholder table supports RACI discussions",
          "Manual editing keeps PM in control before sponsor sign-off"]),
        ("5.5 Work Breakdown Structure (WBS)",
         "Breaks the project into a hierarchical tree with WBS codes, task names, durations, and parent/child relationships. "
         "AI generates from charter; PMs can add, edit, and delete tasks manually.",
         ["Translates deliverables into actionable work packages", "Foundation for scheduling and estimating",
          "Matches manufacturing program structure", "Supports workshop outcomes and expert judgment"]),
        ("5.6 Smart Gantt Chart and Schedule Management",
         "Visual timeline with monthly calendar, task bars, progress, milestones, critical path, and total duration.",
         ["Makes dependencies and timeline visible", "Critical path focus protects end date",
          "Milestones align with phase-gate reviews", "Progress overlay supports weekly status meetings"]),
        ("5.7 Resource Allocation and Department Management",
         "Resources by department with hours, utilization, and expandable employee lists. AI optimization plus manual control.",
         ["Reflects real manufacturing org structure", "Surfaces overallocation early",
          "Supports staffing discussions with functional managers"]),
        ("5.8 Budget and Cost Tracking",
         "Planned vs. actual, variance, forecast, and category breakdown with visual charts.",
         ["Early warning of cost overruns", "Category view explains where money goes",
          "Forecast supports EAC conversations", "Essential for capital and NPI programs"]),
        ("5.9 Inventory and Material Planning",
         "Tracks materials: required vs. available quantities, units, and status. AI-generated or manually managed.",
         ["Prevents delays from missing components", "Links EBOM to procurement readiness",
          "Critical for electronics where chipset/PCB availability drives schedule"]),
        ("5.10 AI Risk Management",
         "Risk register with title, category, probability, impact, score, and mitigation. AI-generated or manual.",
         ["Living register instead of ad hoc emails", "Scoring supports prioritization",
          "Mitigation ensures response plans", "Supports formal risk reviews"]),
        ("5.11 Quality Management System",
         "Defect rate vs. target, quality score, and CAPA records (issue, root cause, action, status).",
         ["Connects delivery to quality outcomes", "CAPA supports ISO-style processes",
          "Coordinates with QA before mass-production sign-off"]),
        ("5.12 Predictive Maintenance Planning",
         "Equipment health score, failure probability, predicted failure date, and recommendations.",
         ["Reduces unplanned downtime during pilot runs", "Aligns schedule with equipment availability",
          "Coordinates with maintenance before critical windows"]),
        ("5.13 Project KPI Dashboard",
         "SPI, CPI, quality metrics, resource utilization, health score, milestone progress, and performance radar chart.",
         ["One screen answers 'How are we doing?'", "SPI/CPI aligns with earned-value practices",
          "Milestone tracking supports phase-gate decisions", "Reduces time compiling status"]),
        ("5.14 AI Project Assistant (Chatbot)",
         "Conversational assistant that answers questions about status, risks, budget, schedule, quality, inventory, "
         "and maintenance. Can modify charter, WBS, inventory, risks, and maintenance through natural language.",
         ["Natural interface for busy PMs", "Speeds data entry during meetings",
          "Onboarding aid for new team members", "Co-pilot that informs and executes updates"]),
        ("5.15 Downloadable Project Report (PowerPoint)",
         "Generates complete editable PowerPoint with overview, charter, WBS, schedule, resources, budget, inventory, "
         "risks, quality, maintenance, and KPIs.",
         ["Ready-made deck for sponsor and gate meetings", "Editable before presenting",
          "Eliminates hours of copy-paste", "Consistent reporting across projects"]),
    ]

    for title, desc, bullets in sections:
        doc.add_heading(title, level=2)
        doc.add_paragraph(desc)
        doc.add_paragraph("Value for project managers:", style="Normal").runs[0].bold = True
        for b in bullets:
            add_bullet(doc, b)

    # 6 Workflow
    doc.add_heading("6. End-to-End Project Manager Workflow", level=1)
    steps = [
        "Login → Dashboard",
        "Create Project (or open existing)",
        "Define Overview (dates, budget, priority)",
        "Generate / edit Charter → Sponsor approval",
        "Generate / edit WBS → Validate with team leads",
        "Generate Gantt → Review critical path & milestones",
        "Plan Resources, Budget, Inventory",
        "Identify & mitigate Risks",
        "Track Quality & Maintenance during execution",
        "Monitor KPIs & health score weekly",
        "Use AI Assistant for quick updates and questions",
        "Export PowerPoint report for reviews",
        "Mark project Completed when closed",
    ]
    for i, s in enumerate(steps, 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(s)

    # 7 Challenge table
    doc.add_heading("7. How the Application Helps Project Managers", level=1)
    add_table(doc, ["PM Challenge", "Application Capability", "Outcome"], [
        ["Planning takes too long", "AI charter, WBS, Gantt, budget, risks", "Faster start, more time for execution"],
        ["Don't know if we're on track", "KPIs, health score, Gantt progress", "Early course correction"],
        ["Budget is slipping", "Budget tab with variance & forecast", "Financial control and credible forecasts"],
        ["Materials will delay us", "Inventory planning with status", "Proactive procurement"],
        ["Risks are in people's heads", "Risk register with mitigations", "Fewer surprises, audit trail"],
        ["Quality issues appear late", "Quality metrics & CAPA", "Shift-left quality focus"],
        ["Machines break during pilot run", "Predictive maintenance view", "Less unplanned downtime"],
        ["Stakeholders want a report", "One-click PowerPoint export", "Professional communication"],
        ["Too busy to update the plan", "Chatbot updates + manual edit tabs", "Plan stays current"],
        ["Too many projects to watch", "Dashboard with health & filters", "Portfolio-level control"],
    ])

    # 8 Differentiators
    doc.add_heading("8. Key Differentiators", level=1)
    for item in [
        "Manufacturing-focused — includes inventory, machines, CAPA, and department-based resources",
        "AI + human control — AI accelerates drafting; PM always has manual override",
        "Unified workspace — one application from charter to KPIs to export",
        "Conversational management — chatbot both advises and updates data",
        "Executive-ready output — PowerPoint reports without separate reporting tools",
    ]:
        add_bullet(doc, item)

    # 9 Conclusion
    doc.add_heading("9. Conclusion", level=1)
    doc.add_paragraph(
        "The AI Project Management Agent is a practical, manufacturing-aware platform that helps project managers "
        "plan faster, see clearer, act earlier, and report better. By combining structured project management "
        "disciplines with AI assistance and conversational control, it reduces administrative burden while improving "
        "the quality and timeliness of decisions."
    )
    doc.add_paragraph(
        "For project managers responsible for complex programs—whether upgrading a production line, developing a new "
        "hardware product, or preparing for mass production—the application provides a single source of truth and an "
        "intelligent partner throughout the project lifecycle."
    )

    footer = doc.add_paragraph()
    footer.add_run("\n\n— End of Report —").italic = True
    note = doc.add_paragraph()
    note.add_run(
        "This report describes application functionality and business value only. "
        "It is intended for project management, stakeholder review, and organizational adoption planning."
    ).italic = True

    doc.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    build()
