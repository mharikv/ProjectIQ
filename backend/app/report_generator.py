"""Generate editable PowerPoint project reports."""

import io
import logging
from datetime import datetime
from typing import Any, Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

logger = logging.getLogger(__name__)

TITLE_COLOR = RGBColor(0x1E, 0x40, 0xAF)
ACCENT_COLOR = RGBColor(0x25, 0x63, 0xEB)


def _fmt_currency(val: Optional[float]) -> str:
    if val is None:
        return "—"
    return f"${val:,.0f}"


def _fmt_date(val) -> str:
    if not val:
        return "—"
    if isinstance(val, str):
        return val
    return val.strftime("%Y-%m-%d") if hasattr(val, "strftime") else str(val)


def _add_title_slide(prs: Presentation, title: str, subtitle: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(8.4), Inches(1.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TITLE_COLOR

    sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(8.4), Inches(0.8))
    sub.text_frame.paragraphs[0].text = subtitle
    sub.text_frame.paragraphs[0].font.size = Pt(18)
    sub.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)


def _add_section_slide(prs: Presentation, section_title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.8), Inches(3), Inches(8.4), Inches(1))
    p = box.text_frame.paragraphs[0]
    p.text = section_title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR


def _add_content_slide(prs: Presentation, title: str, bullets: list[str], max_items: int = 12):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9), Inches(0.7))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(24)
    tp.font.bold = True
    tp.font.color.rgb = TITLE_COLOR

    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(8.8), Inches(5.8))
    tf = body.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets[:max_items]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item if item.startswith("•") else f"• {item}"
        p.font.size = Pt(14)
        p.space_after = Pt(6)
        p.level = 0


def _add_table_slide(prs: Presentation, title: str, headers: list[str], rows: list[list[str]], col_widths=None):
    if not rows:
        _add_content_slide(prs, title, ["No data available for this section."])
        return

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(9), Inches(0.7))
    title_box.text_frame.paragraphs[0].text = title
    title_box.text_frame.paragraphs[0].font.size = Pt(22)
    title_box.text_frame.paragraphs[0].font.bold = True
    title_box.text_frame.paragraphs[0].font.color.rgb = TITLE_COLOR

    cols = len(headers)
    table_rows = min(len(rows) + 1, 14)
    left, top, width, height = Inches(0.5), Inches(1.1), Inches(9), Inches(0.35 * table_rows)
    table = slide.shapes.add_table(table_rows, cols, left, top, width, height).table

    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(11)
            p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_COLOR

    for ri, row in enumerate(rows[: table_rows - 1]):
        for ci, val in enumerate(row[:cols]):
            cell = table.cell(ri + 1, ci)
            cell.text = str(val)[:80]
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)


def _flatten_wbs(nodes: list, depth=0, result=None) -> list[str]:
    result = result or []
    if not isinstance(nodes, list):
        return result
    for node in nodes:
        indent = "  " * depth
        code = node.get("code", "")
        name = node.get("name", "")
        dur = node.get("duration_days", "")
        result.append(f"{indent}{code} {name} ({dur} days)" if dur else f"{indent}{code} {name}")
        if node.get("children"):
            _flatten_wbs(node["children"], depth + 1, result)
    return result


def _normalize_resources(resources: Any) -> list[dict]:
    if not resources:
        return []
    if isinstance(resources, dict):
        items = resources.get("departments", resources.get("allocations", []))
    elif isinstance(resources, list):
        items = resources
    else:
        return []
    return items if isinstance(items, list) else []


def _normalize_risks(risks: Any) -> list[dict]:
    if not risks:
        return []
    if isinstance(risks, dict):
        return risks.get("risks", [])
    if isinstance(risks, list):
        return risks
    return []


def generate_project_report_pptx(project) -> io.BytesIO:
    """Build a complete editable PPTX report for a project model instance."""
    logger.info("Generating PPTX report for project id=%s name='%s'", project.id, project.name)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    report_date = datetime.utcnow().strftime("%B %d, %Y")
    _add_title_slide(
        prs,
        project.name,
        f"Complete Project Report  |  {report_date}\nProjectIQ",
    )

    # Overview
    _add_section_slide(prs, "Project Overview")
    overview = [
        f"Status: {project.status.replace('_', ' ').title()}",
        f"Priority: {project.priority.title()}",
        f"Health Score: {project.health_score}%",
        f"Start Date: {_fmt_date(project.start_date)}",
        f"End Date: {_fmt_date(project.end_date)}",
        f"Planned Budget: {_fmt_currency(project.budget)}",
        f"Actual Cost: {_fmt_currency(project.actual_cost)}",
        "",
        "Description:",
        project.description or "No description provided.",
    ]
    _add_content_slide(prs, "Project Overview", overview)

    # Charter
    charter = project.charter or {}
    if charter and isinstance(charter, dict) and len(charter) > 1:
        _add_section_slide(prs, "Project Charter")
        if charter.get("scope"):
            _add_content_slide(prs, "Scope", [charter["scope"]])
        if charter.get("objectives"):
            _add_content_slide(prs, "Objectives", charter["objectives"])
        if charter.get("deliverables"):
            _add_content_slide(prs, "Deliverables", charter["deliverables"])
        if charter.get("stakeholders"):
            rows = [[s.get("name", ""), s.get("role", ""), s.get("interest", "")] for s in charter["stakeholders"]]
            _add_table_slide(prs, "Stakeholders", ["Name", "Role", "Interest"], rows)
        constraints = (charter.get("constraints") or []) + [
            f"Assumption: {a}" for a in (charter.get("assumptions") or [])
        ]
        if constraints:
            _add_content_slide(prs, "Constraints & Assumptions", constraints)

    # WBS
    wbs = project.wbs or []
    if isinstance(wbs, list) and wbs:
        _add_section_slide(prs, "Work Breakdown Structure")
        wbs_lines = _flatten_wbs(wbs)
        for i in range(0, len(wbs_lines), 12):
            _add_content_slide(prs, "WBS (continued)" if i else "Work Breakdown Structure", wbs_lines[i : i + 12], max_items=12)

    # Gantt / Schedule
    gantt = project.gantt or {}
    tasks = gantt.get("tasks", []) if isinstance(gantt, dict) else []
    if tasks:
        _add_section_slide(prs, "Schedule & Gantt")
        rows = [
            [t.get("wbs_code", t.get("id", "")), t.get("name", ""), t.get("start", ""), t.get("end", ""), f"{t.get('progress', 0)}%"]
            for t in tasks[:20]
        ]
        _add_table_slide(prs, "Task Schedule", ["Code", "Task", "Start", "End", "Progress"], rows)
        if gantt.get("milestones"):
            ms = [f"{m.get('name', '')}: {m.get('date', '')}" for m in gantt["milestones"]]
            _add_content_slide(prs, "Milestones", ms)
        if gantt.get("total_duration_days"):
            _add_content_slide(prs, "Schedule Summary", [f"Total Duration: {gantt['total_duration_days']} days"])

    # Resources
    resources = _normalize_resources(project.resources)
    if resources:
        _add_section_slide(prs, "Resource Allocation")
        dept_rows = []
        for r in resources:
            dept = r.get("department") or r.get("resource", "")
            dept_rows.append([dept, r.get("type", ""), r.get("task", ""), str(r.get("hours", "")), f"{r.get('utilization', '')}%"])
        _add_table_slide(prs, "Departments & Resources", ["Department/Resource", "Type", "Task", "Hours", "Utilization"], dept_rows)
        for r in resources:
            employees = r.get("employees") or []
            if employees:
                dept_name = r.get("department") or r.get("resource", "Department")
                emp_rows = [[e.get("employee_number", ""), e.get("name", ""), e.get("designation", "")] for e in employees]
                _add_table_slide(prs, f"Employees — {dept_name}", ["Emp. No.", "Name", "Designation"], emp_rows)

    # Budget
    budget = project.budget_details or {}
    if budget and isinstance(budget, dict):
        _add_section_slide(prs, "Budget & Cost Tracking")
        summary = [
            f"Planned Budget: {_fmt_currency(budget.get('planned_budget', project.budget))}",
            f"Actual Cost: {_fmt_currency(budget.get('actual_cost', project.actual_cost))}",
            f"Variance: {_fmt_currency(budget.get('variance'))}",
            f"Forecast: {_fmt_currency(budget.get('forecast', {}).get('estimated_final_cost') if isinstance(budget.get('forecast'), dict) else budget.get('forecast'))}",
        ]
        _add_content_slide(prs, "Budget Summary", summary)
        categories = budget.get("categories", [])
        if categories:
            rows = [[c.get("name", ""), _fmt_currency(c.get("planned")), _fmt_currency(c.get("actual")), _fmt_currency(c.get("variance"))] for c in categories]
            _add_table_slide(prs, "Cost Breakdown", ["Category", "Planned", "Actual", "Variance"], rows)

    # Inventory
    inventory = project.inventory or {}
    inv_items = inventory.get("materials", inventory) if isinstance(inventory, dict) else inventory
    if isinstance(inv_items, list) and inv_items:
        _add_section_slide(prs, "Inventory & Materials")
        rows = [
            [m.get("name", ""), str(m.get("required", "")), str(m.get("available", "")), m.get("unit", ""), m.get("status", "")]
            for m in inv_items[:15]
        ]
        _add_table_slide(prs, "Material Status", ["Material", "Required", "Available", "Unit", "Status"], rows)

    # Risks
    risks = _normalize_risks(project.risks)
    if risks:
        _add_section_slide(prs, "Risk Management")
        rows = [
            [r.get("title", ""), r.get("category", ""), r.get("probability", ""), r.get("impact", ""), str(r.get("score", ""))]
            for r in risks[:12]
        ]
        _add_table_slide(prs, "Risk Register", ["Risk", "Category", "Probability", "Impact", "Score"], rows)
        mitigations = [f"{r.get('title', '')}: {r.get('mitigation', '')}" for r in risks if r.get("mitigation")]
        if mitigations:
            _add_content_slide(prs, "Mitigation Strategies", mitigations[:10])

    # Quality
    quality = project.quality or {}
    if quality and isinstance(quality, dict):
        _add_section_slide(prs, "Quality Management")
        q_summary = [
            f"Defect Rate: {quality.get('defect_rate', '—')}%",
            f"Target Defect Rate: {quality.get('target_defect_rate', '—')}%",
            f"Quality Score: {quality.get('quality_score', '—')}/100",
        ]
        _add_content_slide(prs, "Quality Metrics", q_summary)
        capa = quality.get("capa_records", [])
        if capa:
            rows = [[c.get("id", ""), c.get("issue", ""), c.get("status", ""), c.get("due_date", "")] for c in capa]
            _add_table_slide(prs, "CAPA Records", ["ID", "Issue", "Status", "Due Date"], rows)

    # Maintenance
    maintenance = project.maintenance or {}
    machines = maintenance.get("machines", maintenance) if isinstance(maintenance, dict) else maintenance
    if isinstance(machines, list) and machines:
        _add_section_slide(prs, "Predictive Maintenance")
        rows = [
            [m.get("name", ""), f"{m.get('health_score', '')}%", f"{(m.get('failure_probability', 0) or 0) * 100:.0f}%", m.get("predicted_failure_date", ""), m.get("recommendation", "")[:40]]
            for m in machines[:10]
        ]
        _add_table_slide(prs, "Machine Health", ["Machine", "Health", "Failure Risk", "Predicted Date", "Recommendation"], rows)

    # KPIs
    kpis = project.kpis or {}
    if kpis and isinstance(kpis, dict):
        _add_section_slide(prs, "Project KPI Dashboard")
        kpi_lines = []
        sp = kpis.get("schedule_performance", {})
        cp = kpis.get("cost_performance", {})
        qm = kpis.get("quality_metrics", {})
        ru = kpis.get("resource_utilization", {})
        ph = kpis.get("project_health", {})
        if sp:
            kpi_lines.append(f"Schedule Performance Index (SPI): {sp.get('spi', '—')}")
        if cp:
            kpi_lines.append(f"Cost Performance Index (CPI): {cp.get('cpi', '—')}")
        if qm:
            kpi_lines.append(f"Defect Rate: {qm.get('defect_rate', '—')}%  |  First Pass Yield: {qm.get('first_pass_yield', '—')}%")
        if ru:
            kpi_lines.append(f"Resource Utilization: {ru.get('overall', '—')}%")
        if ph:
            kpi_lines.append(f"Project Health: {ph.get('score', '—')}% ({ph.get('status', '')})")
        ms = kpis.get("milestones", {})
        if ms:
            kpi_lines.append(f"Milestones: {ms.get('completed', 0)}/{ms.get('total', 0)} completed")
        _add_content_slide(prs, "Key Performance Indicators", kpi_lines or ["KPI data available in dashboard."])

    # Closing
    _add_title_slide(prs, "End of Report", f"{project.name}  |  Generated {report_date}")

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    logger.info("PPTX report generated for project id=%s (%s slides)", project.id, len(prs.slides))
    return buffer
